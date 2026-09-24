from __future__ import annotations

import io
import zipfile
from collections import Counter
from typing import Any

from openpyxl import load_workbook

from .fingerprint import fingerprint_bytes
from .models import IntakeIssue, IntakeRecord, IntakeResult, IssueSeverity, SourceDescriptor

MAX_XLSX_BYTES = 25 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED = 50 * 1024 * 1024
MAX_XLSX_ROWS = 100_000
MAX_XLSX_RATIO = 50


def ingest_xlsx_bytes(
    data: bytes,
    *,
    source_id: str,
    sheet_name: str | int = 0,
    metadata: dict[str, Any] | None = None,
) -> IntakeResult:
    source = SourceDescriptor(
        source_id,
        "xlsx",
        fingerprint_bytes(data),
        len(data),
        metadata or {},
    )
    issues: list[IntakeIssue] = []

    if len(data) > MAX_XLSX_BYTES:
        issues.append(IntakeIssue("xlsx_too_large", "Workbook exceeds compressed size limit.", IssueSeverity.ERROR))
        return IntakeResult(source, (), tuple(issues), False, {"record_count": 0})

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total_uncompressed = sum(info.file_size for info in archive.infolist())
            ratio = total_uncompressed / max(len(data), 1)
            if total_uncompressed > MAX_XLSX_UNCOMPRESSED or ratio > MAX_XLSX_RATIO:
                issues.append(
                    IntakeIssue(
                        "xlsx_decompression_limit",
                        "Workbook exceeds uncompressed-size/decompression-ratio limit.",
                        IssueSeverity.ERROR,
                    )
                )
                return IntakeResult(
                    source,
                    (),
                    tuple(issues),
                    False,
                    {
                        "record_count": 0,
                        "uncompressed_bytes": total_uncompressed,
                        "compression_ratio": round(ratio, 2),
                    },
                )
    except (zipfile.BadZipFile, OSError) as exc:
        issues.append(IntakeIssue("invalid_xlsx", f"Excel container error: {type(exc).__name__}.", IssueSeverity.ERROR))
        return IntakeResult(source, (), tuple(issues), False, {"record_count": 0})

    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:
        issues.append(IntakeIssue("invalid_xlsx", f"Excel parse error: {type(exc).__name__}.", IssueSeverity.ERROR))
        return IntakeResult(source, (), tuple(issues), False, {"record_count": 0})

    if len(workbook.sheetnames) > 1:
        issues.append(
            IntakeIssue(
                "additional_sheets",
                f"Workbook contains {len(workbook.sheetnames)} sheets; only the selected sheet is ingested.",
                IssueSeverity.WARNING,
            )
        )

    try:
        if isinstance(sheet_name, int):
            worksheet = workbook.worksheets[sheet_name]
        else:
            worksheet = workbook[sheet_name]
    except (IndexError, KeyError):
        issues.append(IntakeIssue("missing_sheet", "Selected worksheet does not exist.", IssueSeverity.ERROR))
        return IntakeResult(source, (), tuple(issues), False, {"record_count": 0})

    header_cells = list(next(worksheet.iter_rows(min_row=1, max_row=1), ()))
    headers = ["" if cell.value is None else str(cell.value).strip() for cell in header_cells]

    if not headers:
        issues.append(IntakeIssue("missing_header", "Worksheet has no header row.", IssueSeverity.ERROR))
    if any(not h for h in headers):
        issues.append(IntakeIssue("blank_header", "Worksheet contains a blank column header.", IssueSeverity.ERROR))
    counts = Counter(headers)
    duplicates = sorted(h for h, count in counts.items() if h and count > 1)
    if duplicates:
        issues.append(
            IntakeIssue(
                "duplicate_headers",
                "Worksheet contains duplicate headers: " + ", ".join(duplicates),
                IssueSeverity.ERROR,
            )
        )

    if any(dim.hidden for dim in worksheet.row_dimensions.values()):
        issues.append(IntakeIssue("hidden_rows", "Worksheet contains hidden rows.", IssueSeverity.WARNING))
    if any(dim.hidden for dim in worksheet.column_dimensions.values()):
        issues.append(IntakeIssue("hidden_columns", "Worksheet contains hidden columns.", IssueSeverity.WARNING))

    records: list[IntakeRecord] = []
    missing = 0
    total = 0
    duplicate_rows = 0
    seen: set[tuple[tuple[str, str], ...]] = set()

    for row_number, cells in enumerate(
        worksheet.iter_rows(min_row=2),
        start=2,
    ):
        if len(records) >= MAX_XLSX_ROWS:
            issues.append(IntakeIssue("too_many_records", "Worksheet row count exceeds limit.", IssueSeverity.ERROR))
            break

        payload: dict[str, Any] = {}
        nonempty = False
        for idx, header in enumerate(headers):
            value = cells[idx].value if idx < len(cells) else None
            if value is not None:
                nonempty = True
            if isinstance(value, str) and value.startswith("="):
                issues.append(
                    IntakeIssue(
                        "formula_cell",
                        f"Formula detected at row {row_number}, column {header or idx+1}.",
                        IssueSeverity.WARNING,
                        header or None,
                    )
                )
            payload[header] = value
            total += 1
            if value is None or (isinstance(value, str) and not value.strip()):
                missing += 1

        if not nonempty:
            continue

        signature = tuple(sorted((str(k), repr(v)) for k, v in payload.items()))
        if signature in seen:
            duplicate_rows += 1
        else:
            seen.add(signature)
        records.append(
            IntakeRecord(
                f"r{len(records)+1}",
                source_id,
                payload,
                row_number,
            )
        )

    if not records:
        issues.append(IntakeIssue("no_data_rows", "Worksheet contains headers but no data rows.", IssueSeverity.ERROR))
    if duplicate_rows:
        issues.append(IntakeIssue("duplicate_rows", f"Detected {duplicate_rows} duplicate row(s).", IssueSeverity.WARNING))

    missing_rate = missing / total if total else 0.0
    if missing:
        issues.append(IntakeIssue("missing_values", f"Blank-cell rate is {missing_rate:.1%}.", IssueSeverity.WARNING))

    valid = not any(issue.severity is IssueSeverity.ERROR for issue in issues)
    stats = {
        "record_count": len(records),
        "column_count": len(headers),
        "missing_cell_rate": round(missing_rate, 6),
        "duplicate_row_count": duplicate_rows,
        "headers": headers,
        "sheet_name": worksheet.title,
        "sheet_count": len(workbook.sheetnames),
    }
    return IntakeResult(source, tuple(records), tuple(issues), valid, stats)
